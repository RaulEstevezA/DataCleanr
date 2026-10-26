"""Excel (.xlsx) import and export. Mirrors web/src/core/xlsx.ts and xlsxSheet.ts.

Every cell becomes text, following shared/test-cases/xlsx.json.
"""

import datetime
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ..core.dataset import Dataset
from ..core.export import ColumnType, parse_excel_date, to_typed_cell
from ..core.jsnumber import js_number_to_string


class EmptySheetError(Exception):
    pass


@dataclass(frozen=True)
class XlsxSheet:
    sheet: str
    data: list[list[object]]


def cell_to_text(value: object) -> str:
    """Converts a cell to text. Dates use ISO order (unambiguous) and only include the time when there is one."""
    if value is None:
        return ""
    # bool before numbers: in Python True is also an int
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int | float):
        return js_number_to_string(float(value)) if isinstance(value, float) else str(value)
    if isinstance(value, datetime.datetime):
        has_time = value.hour or value.minute or value.second
        return value.strftime("%Y-%m-%d %H:%M:%S") if has_time else value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M:%S")
    return str(value)


def read_workbook(path: Path) -> list[XlsxSheet]:
    """Reads every sheet. data_only=True returns the values Excel calculated for formulas."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            XlsxSheet(ws.title, [list(row) for row in ws.iter_rows(values_only=True)]) for ws in workbook.worksheets
        ]
    finally:
        workbook.close()


def sheet_to_dataset(sheet: XlsxSheet, file_name: str) -> Dataset:
    """Builds a dataset from one sheet: first row is the header, fully empty rows are skipped."""
    text_rows = [[cell_to_text(cell) for cell in row] for row in sheet.data]
    text_rows = [row for row in text_rows if any(cell != "" for cell in row)]
    if not text_rows:
        raise EmptySheetError(sheet.sheet)
    header, *body = text_rows
    # Trailing empty header cells come from the sheet's used range, not from real columns
    while header and header[-1] == "" and all(len(row) < len(header) or row[len(header) - 1] == "" for row in body):
        header = header[:-1]
    normalized = [[row[i] if i < len(row) else "" for i in range(len(header))] for row in body]
    return Dataset(
        file_name=file_name,
        format="xlsx",
        sheet=sheet.sheet,
        columns=header,
        rows=normalized,
        row_ids=list(range(len(normalized))),
    )


def write_workbook(sheets: list[tuple[str, Dataset]], types: list[ColumnType], path: Path) -> None:
    """Writes an .xlsx file, one sheet per (name, dataset), with each column converted to its
    chosen type. The header is bold and frozen so it stays visible while scrolling in Excel."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, dataset in sheets:
        _write_sheet(workbook.create_sheet(name), dataset, types)
    workbook.save(path)


def _write_sheet(sheet, dataset: Dataset, types: list[ColumnType]) -> None:
    for c, name in enumerate(dataset.columns, start=1):
        _write_text(sheet.cell(row=1, column=c), name).font = Font(bold=True)

    for r, row in enumerate(dataset.rows, start=2):
        for c, value in enumerate(row, start=1):
            typed = to_typed_cell(value, types[c - 1])
            cell = sheet.cell(row=r, column=c)
            if typed.type == "number" or typed.type == "boolean":
                cell.value = typed.value
            elif typed.type == "date":
                cell.value = parse_excel_date(value)
                cell.number_format = "yyyy-mm-dd hh:mm:ss" if " " in value else "yyyy-mm-dd"
            elif typed.type == "text":
                _write_text(cell, value)

    sheet.freeze_panes = "A2"


def _write_text(cell, value: str):
    cell.value = value
    # openpyxl turns text starting with "=" into a formula; force plain text so "=1+1" stays as written
    cell.data_type = "s"
    return cell

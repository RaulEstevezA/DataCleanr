from openpyxl import load_workbook

from datacleanr.core.dataset import Dataset
from datacleanr.io.csv_file import parse_csv_text
from datacleanr.ui.export_dialog import ExportDialog, write_file

DATASET = Dataset(
    "clientes.csv",
    "csv",
    ["ID", "Importe", "Fecha", "Activo"],
    [["007", "12.5", "2024-07-13", "TRUE"], ["008", "13", "2024-07-14", "false"]],
    [0, 1],
    delimiter=";",
)


def test_suggested_types_and_file_name(window):
    window._i18n.set_language("es")
    dialog = ExportDialog(window._i18n, DATASET, window)
    assert dialog.file_format == "csv"
    assert dialog.file_name.text() == "clientes_clean.csv"
    assert dialog.column_types == ["text", "number", "date", "boolean"]
    dialog.xlsx_button.click()
    assert dialog.file_name.text() == "clientes_clean.xlsx"
    # Forcing ID to number: "007" would lose its zeros, so it stays text and is reported
    dialog._combos[0].setCurrentIndex(1)
    assert dialog.types.item(0, 3).text() == "2 valores se quedarán como texto"


def test_write_csv_and_excel(tmp_path):
    [csv_path] = write_file([(None, "Data", DATASET)], tmp_path / "out.csv", "csv", ",", [])
    # Bytes, not read_text(): universal newlines would hide the CRLF line endings
    text = csv_path.read_bytes().decode("utf-8")
    assert text.startswith("\ufeffID,Importe") and "\r\n" in text
    assert parse_csv_text(text, "out.csv").rows == DATASET.rows

    [xlsx_path] = write_file(
        [(None, "Data", DATASET)], tmp_path / "out.xlsx", "xlsx", ",", ["text", "number", "date", "boolean"]
    )
    sheet = load_workbook(xlsx_path).active
    assert [cell.value for cell in sheet[2]][:2] == ["007", 12.5]
    assert sheet["D2"].value is True


INCOMPLETE = Dataset(
    "clientes.csv",
    "csv",
    ["Nombre", "Email", "Notas"],
    [["Ana", "ana@x.com", "VIP"], ["Juan", "", ""], ["Laura", "laura@x.com", ""]],
    [0, 1, 2],
)


def test_split_into_two_csv_files_and_two_excel_sheets(window, tmp_path):
    window._i18n.set_language("es")
    dialog = ExportDialog(window._i18n, INCOMPLETE, window)
    dialog.split.setChecked(True)
    assert dialog.split_counts.text() == "1 fila completa · 2 filas incompletas"
    dialog.required_buttons[2].click()
    assert dialog.split_counts.text() == "2 filas completas · 1 fila incompleta"
    assert dialog.file_name.text() == "clientes_clean_complete.csv · clientes_clean_incomplete.csv"

    paths = write_file(dialog.parts(), tmp_path / "clientes_clean.csv", "csv", ";", [])
    assert [p.name for p in paths] == ["clientes_clean_complete.csv", "clientes_clean_incomplete.csv"]
    assert parse_csv_text(paths[1].read_text(encoding="utf-8"), "x.csv").rows == [["Juan", "", ""]]

    dialog.xlsx_button.click()
    assert dialog.file_name.text() == "clientes_clean.xlsx"
    [xlsx] = write_file(dialog.parts(), tmp_path / "clientes_clean.xlsx", "xlsx", ",", ["text"] * 3)
    workbook = load_workbook(xlsx)
    assert [(ws.title, ws.max_row - 1) for ws in workbook.worksheets] == [("Completas", 2), ("Incompletas", 1)]


def test_split_skips_an_empty_part(window, tmp_path):
    window._i18n.set_language("es")
    dialog = ExportDialog(window._i18n, DATASET, window)
    dialog.split.setChecked(True)
    assert dialog.split_hint.text().startswith("Una de las dos partes está vacía")
    assert [p.name for p in write_file(dialog.parts(), tmp_path / "x.csv", "csv", ";", [])] == ["x_complete.csv"]
